"""Generate Word and PowerPoint KT deck from markdown outline."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from pptx import Presentation
from pptx.util import Inches as PptxInches, Pt as PptxPt

DOCS_DIR = Path(__file__).resolve().parent
SOURCE = DOCS_DIR / "Fulfillment_Console_QA_KT_SlideDeck_Outline.md"
DOCX_OUT = DOCS_DIR / "Fulfillment_Console_QA_KT.docx"
PPTX_OUT = DOCS_DIR / "Fulfillment_Console_QA_KT.pptx"


def strip_md(text: str) -> str:
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    text = text.replace("❌ ", "").replace("✅ ", "")
    return text.strip()


def parse_slides(content: str) -> list[dict]:
    slides: list[dict] = []
    blocks = re.split(r"\n---\n", content)
    for block in blocks:
        match = re.match(r"## SLIDE (\d+) — (.+)", block.strip())
        if not match:
            continue
        slide_num = int(match.group(1))
        slide_title = match.group(2).strip()
        on_slide_match = re.search(r"\*\*On slide:\*\*\n(.*?)(?=\n\*\*Speaker notes:\*\*|\Z)", block, re.S)
        notes_match = re.search(r"\*\*Speaker notes:\*\*\n(.*)", block, re.S)
        on_slide = on_slide_match.group(1).strip() if on_slide_match else ""
        speaker_notes = notes_match.group(1).strip() if notes_match else ""
        slides.append(
            {
                "num": slide_num,
                "title": slide_title,
                "on_slide": on_slide,
                "speaker_notes": speaker_notes,
            }
        )
    return slides


def on_slide_lines(on_slide: str) -> list[str]:
    lines: list[str] = []
    in_code = False
    for raw in on_slide.splitlines():
        line = raw.rstrip()
        if line.startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            lines.append(line)
            continue
        if line.startswith("|") and "---" in line:
            continue
        if line.startswith("|"):
            cells = [strip_md(c.strip()) for c in line.strip("|").split("|")]
            lines.append(" | ".join(c for c in cells if c))
            continue
        if line.startswith("- ") or re.match(r"^\d+\.\s", line):
            lines.append(strip_md(line.lstrip("- ").strip()))
            continue
        if line.startswith("**") and line.endswith("**"):
            lines.append(strip_md(line))
            continue
        if line.strip():
            lines.append(strip_md(line))
    return lines


def build_docx(slides: list[dict], meta_lines: list[str]) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)

    title = doc.add_heading("Fulfillment Console — QA Knowledge Transfer", 0)
    title.alignment = 0
    for line in meta_lines:
        p = doc.add_paragraph(strip_md(line))
        p.runs[0].font.size = Pt(11)

    doc.add_page_break()

    for slide in slides:
        doc.add_heading(f"Slide {slide['num']}: {slide['title']}", level=1)
        doc.add_heading("On slide", level=2)
        for line in on_slide_lines(slide["on_slide"]):
            doc.add_paragraph(line, style="List Bullet")

        doc.add_heading("Speaker notes", level=2)
        notes = doc.add_paragraph(strip_md(slide["speaker_notes"]))
        notes.runs[0].italic = True
        doc.add_paragraph("")

    doc.save(DOCX_OUT)


def build_pptx(slides: list[dict]) -> None:
    prs = Presentation()
    prs.slide_width = PptxInches(13.333)
    prs.slide_height = PptxInches(7.5)

    title_layout = prs.slide_layouts[0]
    content_layout = prs.slide_layouts[1]

    for slide_data in slides:
        lines = on_slide_lines(slide_data["on_slide"])
        if slide_data["num"] == 1:
            slide = prs.slides.add_slide(title_layout)
            slide.shapes.title.text = "Fulfillment Console — QA Knowledge Transfer"
            subtitle = slide.placeholders[1]
            subtitle.text = "\n".join(lines[:4]) if lines else slide_data["title"]
        else:
            slide = prs.slides.add_slide(content_layout)
            slide.shapes.title.text = f"{slide_data['num']}. {slide_data['title']}"
            body = slide.placeholders[1].text_frame
            body.clear()
            for idx, line in enumerate(lines[:12]):
                if idx == 0:
                    body.text = line
                else:
                    p = body.add_paragraph()
                    p.text = line
                    p.level = 0

        notes_slide = slide.notes_slide
        notes_tf = notes_slide.notes_text_frame
        notes_tf.text = strip_md(slide_data["speaker_notes"])

    prs.save(PPTX_OUT)


def main() -> None:
    content = SOURCE.read_text(encoding="utf-8")
    meta: list[str] = []
    for line in content.splitlines()[3:8]:
        if line.startswith("**"):
            meta.append(line)

    slides = parse_slides(content)
    if not slides:
        raise SystemExit("No slides parsed from markdown source.")

    build_docx(slides, meta)
    build_pptx(slides)
    print(f"Created: {DOCX_OUT}")
    print(f"Created: {PPTX_OUT}")
    print(f"Slides: {len(slides)}")


if __name__ == "__main__":
    main()
