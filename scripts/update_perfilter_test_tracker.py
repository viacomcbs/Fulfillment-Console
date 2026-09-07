#!/usr/bin/env python3
"""
Scan per-filter test classes, Surefire/TestNG reports, and persistent history; update Excel tracker.

Works with:
  - mvn test -DsuiteXmlFile=...
  - IntelliJ / TestNG suite XML runs
  - Manual: python scripts/update_perfilter_test_tracker.py

History is persisted so older runs are kept even when target/surefire-reports is overwritten.
"""

from __future__ import annotations

import json
import re
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parents[1]
PERFILTER_ROOT = (
    PROJECT_ROOT
    / "src/test/java/com/paramount/test/ff/uitests/tests/leftfiltersvalidation/perfilter"
)
HISTORY_JSON = PROJECT_ROOT / "docs/test-tracking/perfilter_execution_history.json"
OUTPUT_XLSX = PROJECT_ROOT / "docs/test-tracking/LeftFilter_PerFilter_TestTracker.xlsx"
ARCHIVE_PATHS_FILE = PROJECT_ROOT / "docs/test-tracking/archive_report_paths.txt"

DEFAULT_ARCHIVE_REPORT_DIRS = [
    Path(r"C:\Users\10747125\OneDrive - LTIMindtree\Desktop\Automation\Fulfillment console\FF_Console_Automation\target\surefire-reports"),
    Path(r"C:\Users\10747125\OneDrive - LTIMindtree\Desktop\Automation\Fulfillment console\FF_Console_Automation\test-output"),
    Path(r"C:\Users\10747125\OneDrive - LTIMindtree\Desktop\Automation\Fulfillment console - Copy\FF_Console_Automation\allure-results"),
]

SKIP_CLASS_SUFFIXES = {
    "LeftFilterOrdersTabBaseTest",
    "LeftFilterLineItemsTabBaseTest",
}

SKIP_METHOD_FRAGMENTS = ("setup", "teardown", "atstart", "before", "after")

TEST_TYPE_MAP = {
    "BasicTest": "Basic",
    "OptionOrderTest": "Option Order",
    "SearchTest": "Search",
    "SelectAllTest": "Select All",
    "TableSyncTest": "Table Sync",
    "ScrollTest": "Scroll",
    "ActiveFiltersTest": "Active Filters",
    "ClearFiltersTest": "Clear Filters",
    "PerFilterValidationTest": "Bulk Validation",
    "CrossTabSwitchValidationTest": "Cross Tab Switch",
}

STATUS_PRIORITY = {"Pass": 3, "Fail": 2, "Skipped": 1}


@dataclass
class TestCaseRow:
    view: str
    filter_name: str
    tc_id: str
    test_type: str
    test_class: str
    test_method: str
    java_file: str
    executed: str = "No"
    latest_status: str = "Not Executed"
    best_status: str = "Not Executed"
    environment: str = ""
    last_run_date: str = ""
    run_count: int = 0
    failure_reason: str = ""


@dataclass
class ExecutionRecord:
    test_class: str
    test_method: str
    status: str
    environment: str
    timestamp: str
    suite_name: str = ""
    failure: str = ""
    source: str = ""


def camel_to_title(name: str) -> str:
    spaced = re.sub(r"([a-z])([A-Z])", r"\1 \2", name)
    return spaced.replace("_", " ").strip().title()


def infer_test_type(class_name: str) -> str:
    for suffix, label in TEST_TYPE_MAP.items():
        if class_name.endswith(suffix):
            return label
    return "Other"


def extract_tc_id(class_name: str) -> str:
    match = re.search(r"(LF_(?:O|LI)_TC\d+)", class_name)
    return match.group(1) if match else class_name


def extract_test_method(java_path: Path) -> str:
    text = java_path.read_text(encoding="utf-8", errors="ignore")
    match = re.search(r"@Test[^)]*\)\s*(?:@[^\n]+\s*)*public\s+void\s+(\w+)\s*\(", text, re.S)
    return match.group(1) if match else ""


def infer_environment(suite_name: str, explicit: str = "") -> str:
    if explicit:
        return explicit.upper()
    upper = suite_name.upper()
    for env in ("PROD", "UAT", "LOCAL", "DEV"):
        if env in upper:
            return env
    return suite_name or "UNKNOWN"


def should_skip_method(method: str) -> bool:
    lower = method.lower()
    if method.startswith("@"):
        return True
    return any(fragment in lower for fragment in SKIP_METHOD_FRAGMENTS)


def testcase_status(testcase: ET.Element) -> str:
    if len(testcase.findall("skipped")) > 0:
        return "Skipped"
    if testcase.find("failure") is not None or testcase.find("error") is not None:
        return "Fail"
    return "Pass"


def failure_message(testcase: ET.Element) -> str:
    failure = testcase.find("failure")
    error = testcase.find("error")
    node = failure if failure is not None else error
    if node is None:
        return ""
    message = node.attrib.get("message") or (node.text or "")
    message = re.sub(r"\s+", " ", message).strip()
    return message[:300]


def load_archive_report_dirs() -> List[Path]:
    dirs: List[Path] = [
        PROJECT_ROOT / "target/surefire-reports",
        PROJECT_ROOT / "test-output",
    ]
    dirs.extend(DEFAULT_ARCHIVE_REPORT_DIRS)

    if ARCHIVE_PATHS_FILE.exists():
        for line in ARCHIVE_PATHS_FILE.read_text(encoding="utf-8").splitlines():
            cleaned = line.strip()
            if cleaned and not cleaned.startswith("#"):
                dirs.append(Path(cleaned))

    unique: List[Path] = []
    seen = set()
    for path in dirs:
        resolved = str(path.resolve()) if path.exists() else str(path)
        if resolved not in seen:
            seen.add(resolved)
            unique.append(path)
    return unique


def load_allure_dirs() -> List[Path]:
    dirs = [PROJECT_ROOT / "allure-results"]
    dirs.extend(
        [
            Path(r"C:\Users\10747125\OneDrive - LTIMindtree\Desktop\Automation\Fulfillment console - Copy\FF_Console_Automation\allure-results"),
            Path(r"C:\Users\10747125\OneDrive - LTIMindtree\Desktop\Automation\Fulfillment console\FF_Console_Automation\allure-results"),
        ]
    )
    unique: List[Path] = []
    seen = set()
    for path in dirs:
        key = str(path)
        if path.exists() and key not in seen:
            seen.add(key)
            unique.append(path)
    return unique


def parse_xml_reports() -> List[ExecutionRecord]:
    records: List[ExecutionRecord] = []
    for report_dir in load_archive_report_dirs():
        if not report_dir.exists():
            continue
        for xml_path in report_dir.rglob("*.xml"):
            if xml_path.name in {"testng-failed.xml", "testng-results.xml", "TEST-TestSuite.xml"}:
                continue
            try:
                root = ET.parse(xml_path).getroot()
            except ET.ParseError:
                continue

            suite_name = xml_path.parent.name
            timestamp = root.attrib.get("timestamp", "")

            for testcase in root.findall(".//testcase"):
                classname = testcase.attrib.get("classname", "")
                method = testcase.attrib.get("name", "")
                if not classname or "perfilter" not in classname:
                    continue
                if should_skip_method(method):
                    continue

                class_short = classname.split(".")[-1]
                if class_short in SKIP_CLASS_SUFFIXES:
                    continue

                records.append(
                    ExecutionRecord(
                        test_class=class_short,
                        test_method=method,
                        status=testcase_status(testcase),
                        environment=infer_environment(suite_name),
                        timestamp=timestamp,
                        suite_name=suite_name,
                        failure=failure_message(testcase),
                        source=f"xml:{report_dir.name}",
                    )
                )
    return records


def parse_allure_results() -> List[ExecutionRecord]:
    records: List[ExecutionRecord] = []
    allure_dirs = load_allure_dirs()
    if not allure_dirs:
        return records

    status_map = {
        "passed": "Pass",
        "failed": "Fail",
        "skipped": "Skipped",
        "broken": "Fail",
    }

    for allure_dir in allure_dirs:
        for json_path in allure_dir.glob("*-result.json"):
            try:
                data = json.loads(json_path.read_text(encoding="utf-8", errors="replace"))
            except (json.JSONDecodeError, OSError):
                continue

            full_name = data.get("fullName", "")
            if "perfilter" not in full_name:
                continue

            parts = full_name.rsplit(".", 1)
            fqcn = parts[0] if parts else full_name
            method = parts[1] if len(parts) > 1 else data.get("name", "")
            class_short = fqcn.split(".")[-1]
            if class_short in SKIP_CLASS_SUFFIXES or should_skip_method(method):
                continue

            labels = {label.get("name"): label.get("value") for label in data.get("labels", [])}
            suite_name = labels.get("parentSuite") or labels.get("suite") or "allure"
            env = labels.get("environment") or labels.get("testEnvironment") or ""

            ts_ms = data.get("stop") or data.get("start")
            timestamp = ""
            if ts_ms:
                timestamp = datetime.fromtimestamp(ts_ms / 1000).strftime("%Y-%m-%dT%H:%M:%S")

            raw_status = data.get("status", "")
            failure = ""
            details = data.get("statusDetails") or {}
            if details.get("message"):
                failure = re.sub(r"\s+", " ", str(details["message"])).strip()[:300]

            records.append(
                ExecutionRecord(
                    test_class=class_short,
                    test_method=method,
                    status=status_map.get(raw_status, raw_status.title()),
                    environment=infer_environment(suite_name, env),
                    timestamp=timestamp,
                    suite_name=suite_name,
                    failure=failure,
                    source="allure",
                )
            )
    return records


def load_history() -> List[ExecutionRecord]:
    if not HISTORY_JSON.exists():
        return []
    try:
        data = json.loads(HISTORY_JSON.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return []

    records: List[ExecutionRecord] = []
    for item in data.get("runs", []):
        records.append(
            ExecutionRecord(
                test_class=item.get("testClass", ""),
                test_method=item.get("testMethod", ""),
                status=item.get("status", ""),
                environment=item.get("environment", ""),
                timestamp=item.get("timestamp", ""),
                suite_name=item.get("suiteName", ""),
                failure=item.get("failure", ""),
                source=item.get("source", "history"),
            )
        )
    return records


def save_history(records: List[ExecutionRecord]) -> None:
    HISTORY_JSON.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "runs": [
            {
                "testClass": r.test_class,
                "testMethod": r.test_method,
                "status": r.status,
                "environment": r.environment,
                "suiteName": r.suite_name,
                "timestamp": r.timestamp,
                "failure": r.failure,
                "source": r.source,
            }
            for r in records
        ]
    }
    HISTORY_JSON.write_text(json.dumps(payload, indent=2), encoding="utf-8")


def record_key(record: ExecutionRecord) -> Tuple[str, str, str, str, str]:
    return (
        record.test_class,
        record.test_method,
        record.status,
        record.timestamp,
        record.suite_name,
    )


def merge_execution_records(*sources: List[ExecutionRecord]) -> List[ExecutionRecord]:
    merged: Dict[Tuple[str, str, str, str, str], ExecutionRecord] = {}
    for source in sources:
        for record in source:
            if not record.test_class:
                continue
            merged[record_key(record)] = record
    return list(merged.values())


def group_by_class(records: List[ExecutionRecord]) -> Dict[str, List[ExecutionRecord]]:
    grouped: Dict[str, List[ExecutionRecord]] = {}
    for record in records:
        grouped.setdefault(record.test_class, []).append(record)
    return grouped


def discover_test_cases() -> List[TestCaseRow]:
    rows: List[TestCaseRow] = []
    if not PERFILTER_ROOT.exists():
        return rows

    for java_path in sorted(PERFILTER_ROOT.rglob("*.java")):
        class_name = java_path.stem
        if class_name in SKIP_CLASS_SUFFIXES:
            continue

        rel = java_path.relative_to(PERFILTER_ROOT)
        parts = rel.parts
        if len(parts) < 3:
            continue

        view_key = parts[0]
        if view_key == "orders":
            view = "Orders View"
        elif view_key == "lineitems":
            view = "Line Items View"
        else:
            continue

        filter_folder = parts[1]
        filter_name = camel_to_title(filter_folder)

        rows.append(
            TestCaseRow(
                view=view,
                filter_name=filter_name,
                tc_id=extract_tc_id(class_name),
                test_type=infer_test_type(class_name),
                test_class=class_name,
                test_method=extract_test_method(java_path),
                java_file=str(java_path.relative_to(PROJECT_ROOT)).replace("\\", "/"),
            )
        )
    return rows


def apply_execution_results(rows: List[TestCaseRow], records: List[ExecutionRecord]) -> None:
    grouped = group_by_class(records)

    def sort_key(rec: ExecutionRecord) -> str:
        return rec.timestamp or ""

    for row in rows:
        class_records = grouped.get(row.test_class, [])
        if not class_records:
            continue

        row.executed = "Yes"
        row.run_count = len(class_records)

        latest = sorted(class_records, key=sort_key)[-1]
        row.latest_status = latest.status
        row.environment = latest.environment
        row.last_run_date = latest.timestamp
        row.failure_reason = latest.failure if latest.status == "Fail" else ""

        best = max(class_records, key=lambda r: STATUS_PRIORITY.get(r.status, 0))
        row.best_status = best.status
        if best.status == "Pass":
            row.failure_reason = ""


def autosize_columns(ws) -> None:
    for col_idx, column_cells in enumerate(ws.columns, start=1):
        max_len = 0
        column = get_column_letter(col_idx)
        for cell in column_cells:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[column].width = min(max_len + 2, 60)


def write_workbook(rows: List[TestCaseRow], all_records: List[ExecutionRecord]) -> None:
    OUTPUT_XLSX.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Per-Filter Tests"

    headers = [
        "View",
        "Filter",
        "TC ID",
        "Test Type",
        "Test Class",
        "Test Method",
        "Java File",
        "Executed",
        "Latest Status",
        "Best Status",
        "Environment",
        "Last Run Date",
        "Run Count",
        "Failure Reason",
    ]

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)

    ws.append(headers)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    status_fills = {
        "Pass": PatternFill("solid", fgColor="C6EFCE"),
        "Fail": PatternFill("solid", fgColor="FFC7CE"),
        "Skipped": PatternFill("solid", fgColor="FFEB9C"),
        "Not Executed": PatternFill("solid", fgColor="EDEDED"),
    }

    for row in rows:
        ws.append(
            [
                row.view,
                row.filter_name,
                row.tc_id,
                row.test_type,
                row.test_class,
                row.test_method,
                row.java_file,
                row.executed,
                row.latest_status,
                row.best_status,
                row.environment,
                row.last_run_date,
                row.run_count,
                row.failure_reason,
            ]
        )
        for col in (9, 10):
            status_cell = ws.cell(row=ws.max_row, column=col)
            fill = status_fills.get(str(status_cell.value), status_fills["Not Executed"])
            status_cell.fill = fill

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    autosize_columns(ws)

    history_ws = wb.create_sheet("Run History")
    history_headers = [
        "Timestamp",
        "Test Class",
        "Test Method",
        "Status",
        "Environment",
        "Suite",
        "Source",
        "Failure Reason",
    ]
    history_ws.append(history_headers)
    for cell in history_ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    for record in sorted(all_records, key=lambda r: r.timestamp or ""):
        history_ws.append(
            [
                record.timestamp,
                record.test_class,
                record.test_method,
                record.status,
                record.environment,
                record.suite_name,
                record.source,
                record.failure,
            ]
        )
    history_ws.freeze_panes = "A2"
    history_ws.auto_filter.ref = history_ws.dimensions
    autosize_columns(history_ws)

    summary = wb.create_sheet("Summary")
    total = len(rows)
    executed = sum(1 for r in rows if r.executed == "Yes")
    passed = sum(1 for r in rows if r.best_status == "Pass")
    failed = sum(1 for r in rows if r.best_status == "Fail")
    skipped = sum(1 for r in rows if r.best_status == "Skipped")
    not_executed = sum(1 for r in rows if r.executed == "No")
    total_invocations = len(all_records)
    unique_classes = len(group_by_class(all_records))

    summary_data = [
        ("Left Filter Per-Filter Test Tracker", ""),
        ("Generated At", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Tracker File", str(OUTPUT_XLSX.relative_to(PROJECT_ROOT)).replace("\\", "/")),
        ("History File", str(HISTORY_JSON.relative_to(PROJECT_ROOT)).replace("\\", "/")),
        ("", ""),
        ("Total Test Cases (unique classes)", total),
        ("Unique Classes Executed", unique_classes),
        ("Total Test Invocations (all runs)", total_invocations),
        ("Executed At Least Once", executed),
        ("Not Executed", not_executed),
        ("Best Status = Pass", passed),
        ("Best Status = Fail", failed),
        ("Best Status = Skipped", skipped),
        ("", ""),
        ("Orders View Total", sum(1 for r in rows if r.view == "Orders View")),
        ("Line Items View Total", sum(1 for r in rows if r.view == "Line Items View")),
        ("", ""),
        ("Auto-update triggers", ""),
        ("Maven", "mvn test (exec-maven-plugin after test phase)"),
        ("TestNG suite XML / IntelliJ", "PerFilterTestTrackerListener on base test classes"),
        ("Manual", "python scripts/update_perfilter_test_tracker.py"),
    ]

    for label, value in summary_data:
        summary.append([label, value])

    summary["A1"].font = Font(bold=True, size=14)
    autosize_columns(summary)

    wb.save(OUTPUT_XLSX)


def main() -> None:
    xml_records = parse_xml_reports()
    allure_records = parse_allure_results()
    history_records = load_history()

    merged = merge_execution_records(history_records, xml_records, allure_records)
    save_history(merged)

    rows = discover_test_cases()
    apply_execution_results(rows, merged)
    write_workbook(rows, merged)

    executed = sum(1 for r in rows if r.executed == "Yes")
    passed = sum(1 for r in rows if r.best_status == "Pass")
    failed = sum(1 for r in rows if r.best_status == "Fail")
    print(f"Updated tracker: {OUTPUT_XLSX}")
    print(f"History: {HISTORY_JSON}")
    print(
        f"Unique classes={len(rows)} executed={executed} pass={passed} fail={failed} "
        f"total_invocations={len(merged)}"
    )


if __name__ == "__main__":
    main()
