#!/usr/bin/env python3
"""Build Platform Weekly QA Status report from Jira MCP JSON outputs."""
import json
import glob
import os
import re
from datetime import datetime
from collections import defaultdict

WEEK_START = datetime(2026, 8, 17).date()
WEEK_END = datetime(2026, 8, 21).date()
AGENT_DIR = r"C:\Users\10747125\.cursor\projects\c-FulfillmentConsole-Fulfillment-Console\agent-tools"
OUT_TSV = r"c:\FulfillmentConsole\Fulfillment-Console\docs\Platform_Weekly_QA_Status_2026-08-17_to_2026-08-21.tsv"
OUT_XLSX = r"c:\FulfillmentConsole\Fulfillment-Console\docs\Platform_Weekly_QA_Status_2026-08-17_to_2026-08-21.xlsx"

JQL_SOURCES = {
    "verified": [
        "187b46d6-bfbd-4e64-a210-bfdd00f4bab9.txt",
        "4f4bcf4d-37a8-4a39-a06a-0752595f5628.txt",
        "verified_archives_week.txt",
    ],
    "sent_reopen": ["ffd4de69-545f-4c69-a90c-7ee730538840.txt"],
    "new": ["d8d68015-02f2-4775-b374-c0d75c97e34c.txt"],
}

EXTRA_KEYS = ["BSD-29515", "BSD-29594", "BSD-29615", "BSD-29856"]
ROSTER = {
    "Sundararajan, Akilandeswari", "Byneni, Alekya", "Kumar, Babloo", "Zurang, Shubham",
    "1, Bhawna", "Suryavanshi, KrishnaKumar", "J, Rahul", "Balasundaram, Ravi",
    "Pawar, Rucha", "Arumugam, Vishnupriya", "K, Dhamodharan", "Dutta, Anjan",
}

NAME_MAP = {
    "1, Bhawna": "Bhawna Kaushik",
    "Suryavanshi, KrishnaKumar": "Krishna Suryavanshi",
    "Sundararajan, Akilandeswari": "Akilandeswari Sundararajan",
    "Kumar, Babloo": "Babloo Kumar",
    "Zurang, Shubham": "Shubham Zurang",
    "J, Rahul": "Rahul J",
    "Balasundaram, Ravi": "Ravi Balasundaram",
    "K, Dhamodharan": "Dhamodharan K",
    "Dutta, Anjan": "Anjan Dutta",
    "Byneni, Alekya": "Alekya Byneni",
    "Pawar, Rucha": "Rucha Pawar",
    "Arumugam, Vishnupriya": "Vishnupriya Arumugam",
}

APP_ORDER = [
    "Fulfillment Console", "MSC Fulfillment Performance Metrics", "Admin Console",
    "Mediahub", "AI Chatbot", "Partner Console", "Media Ingest Console",
    "Order Management Console", "QC Workbench", "Library Manager",
    "Central Repository", "Archives",
]

COMPONENT_TO_APP = {
    "Fulfillment Console": "Fulfillment Console",
    "MSC Fulfillment Performance": "MSC Fulfillment Performance Metrics",
    "Admin Console": "Admin Console",
    "Mediahub": "Mediahub", "Media Hub": "Mediahub",
    "AI Chatbot": "AI Chatbot",
    "Partner Console": "Partner Console",
    "Media Ingest Console": "Media Ingest Console", "MIP": "Media Ingest Console",
    "Order Management Console": "Order Management Console",
    "QC Workbench": "QC Workbench",
    "Library Manager": "Library Manager",
    "Central Repository": "Central Repository",
    "Archives": "Archives",
}

ACTION_PRIORITY = {
    "Verified in PROD": 6, "Verified in UAT": 5, "Verified & closed": 4,
    "Sent to dev": 3, "Reopened": 2, "Newly created": 1,
}


def in_week(ts: str) -> bool:
    d = datetime.strptime(ts[:10], "%Y-%m-%d").date()
    return WEEK_START <= d <= WEEK_END


def fmt_display_name(dn: str) -> str:
    if not dn:
        return ""
    return NAME_MAP.get(dn, f"{dn.split(', ', 1)[1]} {dn.split(', ', 1)[0]}" if ", " in dn else dn)


def fmt_user(user) -> str:
    if not user:
        return ""
    return fmt_display_name(user.get("displayName", ""))


def resolve_app(issue) -> str:
    fields = issue.get("fields", issue)
    summary = fields.get("summary", "")
    comps = [c["name"] for c in fields.get("components", [])]
    primary = comps[0] if comps else ""
    if primary == "Order Management Console" and re.search(r"\b(MIC|OC-MIC)\b", summary, re.I):
        return "Media Ingest Console"
    return COMPONENT_TO_APP.get(primary, primary or "Unknown")


def classify_change(field: str, to_str: str):
    if field in ("UAT QA Status", "customfield_16437"):
        if to_str == "Working":
            return "Verified in UAT", "UAT", False
        if to_str == "-Not Working":
            return "Sent to dev", "UAT", True
    if field in ("PROD QA Status", "customfield_16438"):
        if to_str == "Working":
            return "Verified in PROD", "PROD", False
        if to_str == "-Not Working":
            return "Sent to dev", "PROD", True
    if field == "status":
        if to_str in ("Complete", "Done"):
            return "Verified & closed", "Dev", False
        if to_str == "In Progress (DEV)":
            return "Sent to dev", "", True
        if to_str == "Reopened":
            return "Reopened", "", True
    return None


def env_from_description(desc):
    if not desc:
        return ""
    text = desc if isinstance(desc, str) else json.dumps(desc)
    for pat in [r"Environment\s*[:\-]\s*(PROD/UAT|PROD|UAT|Dev)", r"\[(LM2\.0|LMv2)\s+(UAT|PROD|DEV)", r"\b(PROD/UAT|PROD|UAT|Dev)\b"]:
        m = re.search(pat, text, re.I)
        if m:
            g = m.group(m.lastindex)
            return {"dev": "Dev", "prod": "PROD", "uat": "UAT"}.get(g.lower(), g)
    return ""


def load_jql_issues():
    issues = {}
    jql_action = {}
    for qtype, files in JQL_SOURCES.items():
        default_action = {
            "verified": ("Verified & closed", "Dev", False),
            "sent_reopen": None,
            "new": ("Newly created", "", False),
        }[qtype]
        for fname in files:
            path = os.path.join(AGENT_DIR, fname)
            if not os.path.exists(path):
                continue
            data = json.load(open(path, encoding="utf-8"))
            for issue in data.get("issues", []):
                key = issue["key"]
                issues[key] = issue
                if qtype == "sent_reopen":
                    st = issue["fields"]["status"]["name"]
                    jql_action[key] = ("Reopened", "", True) if st == "Reopened" else ("Sent to dev", "", True)
                else:
                    jql_action[key] = default_action
    return issues, jql_action


def changelog_freshness(d):
    histories = d.get("changelog", {}).get("histories", [])
    if not histories:
        return ""
    return max(h.get("created", "") for h in histories)


def load_changelog_issues():
    issues = {}
    for path in glob.glob(os.path.join(AGENT_DIR, "*.txt")):
        try:
            data = json.load(open(path, encoding="utf-8"))
        except Exception:
            continue
        if isinstance(data, dict) and "key" in data and "changelog" in data:
            key = data["key"]
            existing = issues.get(key)
            if not existing or changelog_freshness(data) > changelog_freshness(existing):
                issues[key] = data
    return issues


def scan_changelog_entries(issue_data):
    """All qualifying QA roster changelog hits in the week."""
    entries = []
    for h in issue_data.get("changelog", {}).get("histories", []):
        author = h.get("author", {}).get("displayName", "")
        if author not in ROSTER:
            continue
        created = h.get("created", "")
        if not in_week(created):
            continue
        for item in h.get("items", []):
            result = classify_change(item.get("field", ""), item.get("toString", ""))
            if not result:
                continue
            action, env, notes_req = result
            entries.append({
                "action": action, "env": env, "notes_req": notes_req,
                "created": created, "author": author,
                "priority": ACTION_PRIORITY.get(action, 0),
            })
    return entries


def pick_latest_action(entries):
    if not entries:
        return None
    # Prefer UAT/PROD over Dev Complete when QA progresses environments (PROD often set just before Complete)
    best = max(entries, key=lambda e: (e["priority"], e["created"]))
    return best["action"], best["env"], best["notes_req"], best["author"]


def resolve_name(fields, entries, action):
    itype = fields.get("issuetype", {}).get("name", "")
    reporter = fields.get("reporter")

    if action == "Newly created" and itype == "Bug":
        non_create = [e for e in entries if e["action"] != "Newly created"]
        if not non_create:
            return fmt_user(reporter)

    if entries:
        best = max(entries, key=lambda e: (e["priority"], e["created"]))
        return fmt_display_name(best["author"])

    return fmt_user(reporter)


def merge_issue(key, jql_issue, cl_issue):
    merged = dict(jql_issue) if jql_issue else {}
    if cl_issue:
        merged["changelog"] = cl_issue.get("changelog", {})
        cf = cl_issue.get("fields", {})
        mf = merged.setdefault("fields", {})
        for f in ("customfield_10224", "description", "summary", "status", "issuetype", "components", "reporter", "created"):
            if cf.get(f) is not None:
                mf[f] = cf[f]
        if not merged.get("key"):
            merged["key"] = cl_issue.get("key", key)
    return merged


def build_rows(jql_issues, jql_action, cl_issues):
    rows = {}

    for key in EXTRA_KEYS:
        if key in cl_issues and key not in jql_issues:
            jql_issues[key] = cl_issues[key]
            jql_action[key] = ("Verified & closed", "Dev", False)

    all_keys = set(jql_issues.keys())
    for key, cl in cl_issues.items():
        if scan_changelog_entries(cl):
            all_keys.add(key)
            if key not in jql_issues:
                jql_issues[key] = cl

    for key in sorted(all_keys):
        issue = merge_issue(key, jql_issues.get(key), cl_issues.get(key))
        fields = issue.get("fields", {})
        if not fields:
            continue

        entries = scan_changelog_entries(issue)
        action_info = pick_latest_action(entries)
        if not action_info and key in jql_action:
            fa = jql_action[key]
            action_info = (fa[0], fa[1], fa[2], fields.get("reporter", {}).get("displayName", ""))
        if not action_info:
            continue

        action, env, notes_req, _author = action_info
        name = resolve_name(fields, entries, action)
        if not env:
            env = env_from_description(fields.get("description", ""))
        if action in ("Sent to dev", "Reopened") and notes_req:
            note = "QA sent back to dev during week" if action == "Sent to dev" else "QA reopened during week"
        else:
            note = ""

        rows[key] = {
            "application": resolve_app(issue), "name": name, "key": key,
            "status": fields.get("status", {}).get("name", ""),
            "action": action, "summary": fields.get("summary", ""),
            "notes": note, "environment": env,
            "type": fields.get("issuetype", {}).get("name", ""),
        }

    return list(rows.values())


def write_tsv(rows):
    rows.sort(key=lambda r: (APP_ORDER.index(r["application"]) if r["application"] in APP_ORDER else 99, r["key"]))
    os.makedirs(os.path.dirname(OUT_TSV), exist_ok=True)
    with open(OUT_TSV, "w", encoding="utf-8", newline="") as f:
        f.write("Application\tName\tJira ID\tStatus\tQA action this week\tSummary\tNotes\tEnvironment\tType\n")
        for r in rows:
            f.write("\t".join([
                r["application"], r["name"], r["key"], r["status"], r["action"],
                r["summary"].replace("\t", " "), r["notes"], r["environment"], r["type"],
            ]) + "\n")


def write_xlsx(rows):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "QA Status"
    ws.append(["Application", "Name", "Jira ID", "Status", "QA action this week", "Summary", "Notes", "Environment", "Type"])
    for r in sorted(rows, key=lambda r: (APP_ORDER.index(r["application"]) if r["application"] in APP_ORDER else 99, r["key"])):
        ws.append([r["application"], r["name"], r["key"], r["status"], r["action"], r["summary"], r["notes"], r["environment"], r["type"]])
        c = ws.cell(row=ws.max_row, column=3)
        c.hyperlink = f"https://paramount.atlassian.net/browse/{r['key']}"
        c.font = Font(color="0563C1", underline="single")

    totals = wb.create_sheet("Totals")
    totals.append(["Application", "Verified Dev", "Verified UAT", "Verified PROD", "Sent to dev", "Reopened", "Newly created", "Total"])
    by_app = defaultdict(lambda: defaultdict(int))
    for r in rows:
        m = {"Verified & closed": "dev", "Verified in UAT": "uat", "Verified in PROD": "prod",
             "Sent to dev": "sent", "Reopened": "reopened", "Newly created": "new"}
        by_app[r["application"]][m.get(r["action"], "other")] += 1
        by_app[r["application"]]["total"] += 1
    grand = defaultdict(int)
    for app in APP_ORDER:
        d = by_app[app]
        if not d["total"]:
            continue
        totals.append([app, d["dev"], d["uat"], d["prod"], d["sent"], d["reopened"], d["new"], d["total"]])
        for k, v in d.items():
            grand[k] += v
    totals.append(["GRAND TOTAL", grand["dev"], grand["uat"], grand["prod"], grand["sent"], grand["reopened"], grand["new"], grand["total"]])
    wb.save(OUT_XLSX)


if __name__ == "__main__":
    jql_issues, jql_action = load_jql_issues()
    cl_issues = load_changelog_issues()
    rows = build_rows(jql_issues, jql_action, cl_issues)
    write_tsv(rows)
    write_xlsx(rows)
    print(f"Rows: {len(rows)}")
    by_app = defaultdict(lambda: defaultdict(int))
    for r in rows:
        m = {"Verified & closed": "dev", "Verified in UAT": "uat", "Verified in PROD": "prod",
             "Sent to dev": "sent", "Reopened": "reopened", "Newly created": "new"}
        by_app[r["application"]][m.get(r["action"], "x")] += 1
    for app in APP_ORDER:
        d = by_app[app]
        if any(d.values()):
            print(f"  {app}: {dict(d)}")
    for key in ("BSD-29582", "BSD-27316"):
        hit = next((r for r in rows if r["key"] == key), None)
        if hit:
            print(f"  {key}: {hit['action']} | {hit['name']} | {hit['environment']}")
